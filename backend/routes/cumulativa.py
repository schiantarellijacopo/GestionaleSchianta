"""Cumulativa Assicurati — elenco cumulativo di persone assicurate su una polizza
non-RCA (es. 'Ritiro Patente'). Struttura ispirata all'Excel del cliente:
NUMERO · COGNOME · NOME · DITTA · ENTRATA · USCITA.

Ogni voce è un documento in ``assicurati_cumulativi`` collegato alla polizza.
"""
from __future__ import annotations

import csv
import io
import re
import uuid
from datetime import datetime, date, timezone
from typing import Optional

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from pydantic import BaseModel

from auth import current_user, require_user
from database import db
from db_models import _now_iso
from shared import log_attivita

router = APIRouter()


def _parse_date(value) -> Optional[str]:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    s = str(value).strip()
    if not s or s == "-":
        return None
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def _stato_da_uscita(data_uscita: Optional[str]) -> str:
    if not data_uscita:
        return "attivo"
    try:
        return "uscito" if date.fromisoformat(data_uscita[:10]) <= date.today() else "attivo"
    except (TypeError, ValueError):
        return "uscito"


# ============================================================================
# LISTA
# ============================================================================
@router.get("/polizze/{pid}/assicurati")
async def list_assicurati(pid: str, user=Depends(current_user)) -> list[dict]:
    docs = await db.assicurati_cumulativi.find(
        {"polizza_id": pid}, {"_id": 0},
    ).sort("numero", 1).to_list(5000)
    return docs


@router.get("/polizze/{pid}/assicurati/stats")
async def stats_assicurati(pid: str, user=Depends(current_user)) -> dict:
    docs = await db.assicurati_cumulativi.find({"polizza_id": pid}, {"_id": 0, "stato": 1}).to_list(5000)
    return {
        "totale": len(docs),
        "attivi": sum(1 for d in docs if d.get("stato") == "attivo"),
        "usciti": sum(1 for d in docs if d.get("stato") == "uscito"),
    }


# ============================================================================
# CRUD
# ============================================================================
class AssicuratoBody(BaseModel):
    cognome: str
    nome: Optional[str] = None
    ditta: Optional[str] = None
    data_entrata: Optional[str] = None
    data_uscita: Optional[str] = None
    codice_fiscale: Optional[str] = None
    anagrafica_id: Optional[str] = None
    note: Optional[str] = None


async def _next_numero(pid: str) -> int:
    last = await db.assicurati_cumulativi.find(
        {"polizza_id": pid}, {"_id": 0, "numero": 1},
    ).sort("numero", -1).limit(1).to_list(1)
    return (last[0]["numero"] + 1) if last else 1


@router.post("/polizze/{pid}/assicurati", status_code=201)
async def create_assicurato(
    pid: str, body: AssicuratoBody,
    user=Depends(require_user("admin", "collaboratore", "dipendente")),
) -> dict:
    pol = await db.polizze.find_one({"id": pid}, {"_id": 0, "id": 1})
    if not pol:
        raise HTTPException(404, "Polizza non trovata")
    doc = {
        "id": str(uuid.uuid4()),
        "polizza_id": pid,
        "numero": await _next_numero(pid),
        "cognome": (body.cognome or "").strip().upper(),
        "nome": (body.nome or "").strip().upper() or None,
        "ditta": (body.ditta or "").strip() or None,
        "data_entrata": _parse_date(body.data_entrata),
        "data_uscita": _parse_date(body.data_uscita),
        "codice_fiscale": (body.codice_fiscale or "").strip().upper() or None,
        "anagrafica_id": body.anagrafica_id,
        "note": body.note,
        "stato": _stato_da_uscita(_parse_date(body.data_uscita)),
        "created_at": _now_iso(),
        "updated_at": _now_iso(),
    }
    await db.assicurati_cumulativi.insert_one(doc)
    await db.polizze.update_one({"id": pid}, {"$set": {"is_cumulativa_assicurati": True}})
    await log_attivita(user, "create", "assicurato_cumulativo", doc["id"])
    doc.pop("_id", None)
    return doc


@router.put("/assicurati/{aid}")
async def update_assicurato(
    aid: str, body: AssicuratoBody,
    user=Depends(require_user("admin", "collaboratore", "dipendente")),
) -> dict:
    upd = {
        "cognome": (body.cognome or "").strip().upper(),
        "nome": (body.nome or "").strip().upper() or None,
        "ditta": (body.ditta or "").strip() or None,
        "data_entrata": _parse_date(body.data_entrata),
        "data_uscita": _parse_date(body.data_uscita),
        "codice_fiscale": (body.codice_fiscale or "").strip().upper() or None,
        "anagrafica_id": body.anagrafica_id,
        "note": body.note,
        "stato": _stato_da_uscita(_parse_date(body.data_uscita)),
        "updated_at": _now_iso(),
    }
    res = await db.assicurati_cumulativi.update_one({"id": aid}, {"$set": upd})
    if res.matched_count == 0:
        raise HTTPException(404, "Assicurato non trovato")
    return await db.assicurati_cumulativi.find_one({"id": aid}, {"_id": 0})


@router.delete("/assicurati/{aid}")
async def delete_assicurato(
    aid: str, user=Depends(require_user("admin", "collaboratore", "dipendente")),
) -> dict:
    res = await db.assicurati_cumulativi.delete_one({"id": aid})
    if res.deleted_count == 0:
        raise HTTPException(404, "Assicurato non trovato")
    return {"ok": True}


# ============================================================================
# IMPORT EXCEL/CSV
# ============================================================================
_HEADER_MAP = {
    "cognome": "cognome",
    "nome": "nome",
    "ditta": "ditta",
    "azienda": "ditta",
    "societa": "ditta",
    "entrata": "data_entrata",
    "data entrata": "data_entrata",
    "inizio": "data_entrata",
    "uscita": "data_uscita",
    "data uscita": "data_uscita",
    "fine": "data_uscita",
    "codice fiscale": "codice_fiscale",
    "cf": "codice_fiscale",
}


def _norm(h) -> str:
    return re.sub(r"\s+", " ", str(h or "").strip().lower())


def _find_header_idx(rows) -> int:
    """Trova la riga header: quella che contiene più parole chiave note
    (cognome/nome/ditta/entrata/uscita/numero). Fallback: prima riga con >=2 celle."""
    keys = {"cognome", "nome", "ditta", "entrata", "uscita", "numero", "azienda", "codice fiscale"}
    best_idx, best_score = None, 0
    for i, r in enumerate(rows[:15]):
        cells = [_norm(c) for c in r]
        score = sum(1 for c in cells if any(k == c or k in c for k in keys))
        if score > best_score:
            best_score, best_idx = score, i
    if best_idx is not None and best_score >= 1:
        return best_idx
    for i, r in enumerate(rows):
        if sum(1 for c in r if c not in (None, "")) >= 2:
            return i
    return 0


def _parse_rows(file_bytes: bytes, filename: str) -> tuple[list[str], list[list]]:
    name = (filename or "").lower()
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return [], []
        header_idx = _find_header_idx(rows)
        header = [_norm(c) for c in rows[header_idx]]
        data = [list(r) for r in rows[header_idx + 1:] if any(c not in (None, "") for c in r)]
        return header, data
    text = file_bytes.decode("utf-8-sig", errors="ignore")
    delim = ";" if text[:2048].count(";") > text[:2048].count(",") else ","
    rows = list(csv.reader(io.StringIO(text), delimiter=delim))
    if not rows:
        return [], []
    header_idx = _find_header_idx(rows)
    header = [_norm(c) for c in rows[header_idx]]
    data = [r for r in rows[header_idx + 1:] if any((c or "").strip() for c in r)]
    return header, data


@router.post("/polizze/{pid}/assicurati/import")
async def import_assicurati(
    pid: str,
    file: UploadFile = File(...),
    sostituisci: bool = Form(False),
    user=Depends(require_user("admin", "collaboratore", "dipendente")),
) -> dict:
    pol = await db.polizze.find_one({"id": pid}, {"_id": 0, "id": 1})
    if not pol:
        raise HTTPException(404, "Polizza non trovata")
    content = await file.read()
    header, data = _parse_rows(content, file.filename or "")
    if not header:
        raise HTTPException(400, "File vuoto o non leggibile")

    # Mappa colonna → campo
    col_field = {}
    for idx, h in enumerate(header):
        key = _norm(h)
        if key in _HEADER_MAP:
            col_field[idx] = _HEADER_MAP[key]
        else:
            for hint, field in _HEADER_MAP.items():
                if hint in key:
                    col_field[idx] = field
                    break
    if "cognome" not in col_field.values():
        raise HTTPException(400, f"Colonna COGNOME non trovata. Header rilevati: {header}")

    if sostituisci:
        await db.assicurati_cumulativi.delete_many({"polizza_id": pid})

    numero = await _next_numero(pid)
    creati = 0
    for row in data:
        rec = {}
        for idx, field in col_field.items():
            raw = row[idx] if idx < len(row) else None
            if field in ("data_entrata", "data_uscita"):
                rec[field] = _parse_date(raw)
            else:
                rec[field] = (str(raw).strip() if raw not in (None, "") else None)
        cognome = (rec.get("cognome") or "").strip().upper()
        if not cognome:
            continue
        doc = {
            "id": str(uuid.uuid4()),
            "polizza_id": pid,
            "numero": numero,
            "cognome": cognome,
            "nome": (rec.get("nome") or "").strip().upper() or None,
            "ditta": (rec.get("ditta") or "").strip() or None,
            "data_entrata": rec.get("data_entrata"),
            "data_uscita": rec.get("data_uscita"),
            "codice_fiscale": (rec.get("codice_fiscale") or "").strip().upper() or None,
            "stato": _stato_da_uscita(rec.get("data_uscita")),
            "created_at": _now_iso(),
            "updated_at": _now_iso(),
        }
        await db.assicurati_cumulativi.insert_one(doc)
        numero += 1
        creati += 1

    await db.polizze.update_one({"id": pid}, {"$set": {"is_cumulativa_assicurati": True}})
    await log_attivita(user, "import", "assicurati_cumulativi", pid)
    return {"ok": True, "creati": creati, "sostituiti": sostituisci, "colonne_mappate": list(col_field.values())}
