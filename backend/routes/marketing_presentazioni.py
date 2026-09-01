"""Studio Marketing — Presentazioni AI (PPTX editabile + PDF).

- Modalità "argomento": l'AI crea struttura e testi di tutte le slide
- Modalità "punti": l'utente fornisce i punti chiave, l'AI li impagina
- Stile "ai": sfondi fotografici Nano Banana 16:9 | Stile "brand": template pulito con colori agenzia
- Output: PPTX (testi modificabili) + PDF + anteprime JPEG
"""
from __future__ import annotations

import asyncio
import io
import logging
import os

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response

from auth import require_user
from database import db
from db_models import _now_iso, _uid
import storage as obj_storage
from routes.marketing_studio import (
    STAFF, FONT_BOLD, FONT_REG, APP_NAME,
    _llm_text, _parse_json, _genera_sfondo, _brand_name, _logo_img,
    _cover_resize, _wrap, _applica_logo,
)

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/marketing")

W, H = 1920, 1080
MARGIN = 116
PX_IN = 144.0  # 1920px ↔ 13.333in

DARK = (15, 23, 42)
ACCENT = (192, 38, 211)
ACCENT_LIGHT = (232, 121, 249)


def _theme(stile: str, tipo: str) -> dict:
    if stile == "brand" and tipo == "contenuto":
        return {"title": DARK, "body": (51, 65, 85), "muted": (100, 116, 139), "accent": ACCENT}
    return {"title": (255, 255, 255), "body": (226, 232, 240),
            "muted": (148, 163, 184), "accent": ACCENT_LIGHT}


def _bg_slide(stile: str, tipo: str, bg_bytes: bytes | None, logo) -> "object":
    """Sfondo senza testi (usato anche come background PPTX)."""
    from PIL import Image, ImageDraw
    if stile == "ai" and bg_bytes:
        img = _cover_resize(Image.open(io.BytesIO(bg_bytes)).convert("RGB"), W, H)
        grad = Image.new("L", (W, 1), 0)
        for x in range(W):
            p = max(0.0, 1 - x / (W * 0.78))
            grad.putpixel((x, 0), int(215 * (p ** 1.1)))
        overlay = Image.new("RGB", (W, H), (8, 14, 24))
        img = Image.composite(overlay, img, grad.resize((W, H)))
    elif stile == "brand" and tipo == "contenuto":
        img = Image.new("RGB", (W, H), (250, 250, 252))
        d = ImageDraw.Draw(img)
        d.rectangle((0, 0, 18, H), fill=ACCENT)
        d.rectangle((0, H - 8, W, H), fill=(226, 232, 240))
    else:
        img = Image.new("RGB", (W, H), DARK)
        d = ImageDraw.Draw(img)
        d.rectangle((0, 0, 18, H), fill=ACCENT)
        d.ellipse((W - 560, H - 560, W + 260, H + 260), fill=(30, 41, 59))
    return _applica_logo(img, logo, w=W, h=H)


def _draw_slide(bg_img, sl: dict, *, stile: str, n: int, tot: int, brand: str) -> bytes:
    """Sfondo + testi → JPEG completo (per PDF e anteprime)."""
    from PIL import ImageDraw, ImageFont
    img = bg_img.copy()
    d = ImageDraw.Draw(img)
    tipo = sl["tipo"]
    th = _theme(stile, tipo)
    f_small = ImageFont.truetype(FONT_BOLD, 26)
    max_w = W - 2 * MARGIN

    if tipo == "copertina":
        f_title = ImageFont.truetype(FONT_BOLD, 96)
        f_sub = ImageFont.truetype(FONT_REG, 44)
        title_lines = _wrap(d, (sl.get("titolo") or "").upper(), f_title, max_w - 200)
        sub_lines = _wrap(d, sl.get("sottotitolo") or "", f_sub, max_w - 300)
        total_h = len(title_lines) * 112 + (len(sub_lines) * 58 + 30 if sub_lines else 0)
        y = (H - total_h) // 2
        d.rectangle((MARGIN, y - 42, MARGIN + 130, y - 30), fill=th["accent"])
        for ln in title_lines:
            d.text((MARGIN, y), ln, font=f_title, fill=th["title"])
            y += 112
        y += 30
        for ln in sub_lines:
            d.text((MARGIN, y), ln, font=f_sub, fill=th["body"])
            y += 58
        d.text((MARGIN, H - MARGIN + 30), brand.upper(), font=f_small, fill=th["muted"])
    elif tipo == "contatti":
        f_title = ImageFont.truetype(FONT_BOLD, 80)
        f_line = ImageFont.truetype(ImageFont.truetype(FONT_REG, 44).path, 44)
        lines = sl.get("punti") or []
        total_h = 100 + 46 + len(lines) * 66
        y = (H - total_h) // 2
        for ln in _wrap(d, (sl.get("titolo") or "Parliamone insieme").upper(), f_title, max_w):
            tw = d.textlength(ln, font=f_title)
            d.text(((W - tw) // 2, y), ln, font=f_title, fill=th["title"])
            y += 100
        y += 46
        for ln in lines:
            tw = d.textlength(ln, font=f_line)
            d.text(((W - tw) // 2, y), ln, font=f_line, fill=th["body"])
            y += 66
        bw = d.textlength(brand.upper(), font=f_small)
        d.text(((W - bw) // 2, H - MARGIN + 30), brand.upper(), font=f_small, fill=th["muted"])
    else:  # contenuto
        f_title = ImageFont.truetype(FONT_BOLD, 62)
        f_body = ImageFont.truetype(FONT_REG, 40)
        if tot > 1:
            d.text((W - MARGIN - d.textlength(f"{n} / {tot}", font=f_small), MARGIN - 10),
                   f"{n} / {tot}", font=f_small, fill=th["muted"])
        y = MARGIN
        for ln in _wrap(d, sl.get("titolo") or "", f_title, max_w - 240):
            d.text((MARGIN, y), ln, font=f_title, fill=th["title"])
            y += 78
        d.rectangle((MARGIN, y + 14, MARGIN + 110, y + 24), fill=th["accent"])
        y += 74
        for punto in (sl.get("punti") or [])[:8]:
            wrapped = _wrap(d, punto, f_body, max_w - 70)
            d.ellipse((MARGIN + 6, y + 16, MARGIN + 24, y + 34), fill=th["accent"])
            for j, ln in enumerate(wrapped):
                d.text((MARGIN + 52, y), ln, font=f_body, fill=th["body"])
                y += 54
            y += 22
        d.text((MARGIN, H - MARGIN + 30), brand.upper(), font=f_small, fill=th["muted"])

    out = io.BytesIO()
    img.convert("RGB").save(out, format="JPEG", quality=90)
    return out.getvalue()


def _pptx_build(slides: list[dict], bg_jpegs: list[bytes], stile: str, brand: str) -> bytes:
    """PPTX 16:9 con sfondi immagine e caselle di testo modificabili."""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    def _in(px):
        return Inches(px / PX_IN)

    prs = Presentation()
    prs.slide_width, prs.slide_height = _in(W), _in(H)
    blank = prs.slide_layouts[6]

    def _box(slide, x, y, w, h, lines, *, size, bold=False, color=(255, 255, 255),
             align=PP_ALIGN.LEFT, bullet=False):
        tb = slide.shapes.add_textbox(_in(x), _in(y), _in(w), _in(h))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, txt in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run()
            r.text = f"•  {txt}" if bullet else txt
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = RGBColor(*color)
            if bullet:
                p.space_after = Pt(10)
        return tb

    tot = len(slides)
    for i, (sl, bg) in enumerate(zip(slides, bg_jpegs), start=1):
        tipo = sl["tipo"]
        th = _theme(stile, tipo)
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(io.BytesIO(bg), 0, 0, width=prs.slide_width, height=prs.slide_height)
        if tipo == "copertina":
            _box(slide, MARGIN, 330, W - 2 * MARGIN - 120, 320, [sl.get("titolo") or ""],
                 size=44, bold=True, color=th["title"])
            if sl.get("sottotitolo"):
                _box(slide, MARGIN, 660, W - 2 * MARGIN - 240, 160, [sl["sottotitolo"]],
                     size=21, color=th["body"])
            _box(slide, MARGIN, H - 100, 900, 60, [brand.upper()], size=12, color=th["muted"])
        elif tipo == "contatti":
            _box(slide, MARGIN, 300, W - 2 * MARGIN, 130,
                 [sl.get("titolo") or "Parliamone insieme"],
                 size=38, bold=True, color=th["title"], align=PP_ALIGN.CENTER)
            _box(slide, MARGIN, 470, W - 2 * MARGIN, 380, sl.get("punti") or [],
                 size=20, color=th["body"], align=PP_ALIGN.CENTER)
        else:
            _box(slide, MARGIN, MARGIN - 16, W - 2 * MARGIN - 200, 150, [sl.get("titolo") or ""],
                 size=29, bold=True, color=th["title"])
            _box(slide, MARGIN, 300, W - 2 * MARGIN, H - 300 - MARGIN,
                 (sl.get("punti") or [])[:8], size=18, color=th["body"], bullet=True)
            if tot > 1:
                _box(slide, W - MARGIN - 200, MARGIN - 20, 200, 50, [f"{i} / {tot}"],
                     size=12, color=th["muted"], align=PP_ALIGN.RIGHT)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


async def _linee_contatti() -> list[str]:
    cfg = await db.azienda_config.find_one(
        {}, {"_id": 0, "telefono": 1, "email": 1, "sito_web": 1,
             "indirizzo": 1, "comune": 1, "provincia": 1}) or {}
    lines = []
    if cfg.get("telefono"):
        lines.append(f"Tel. {cfg['telefono']}")
    if cfg.get("email"):
        lines.append(cfg["email"])
    if cfg.get("sito_web"):
        lines.append(cfg["sito_web"])
    addr = " ".join(filter(None, [cfg.get("indirizzo"), "-" if cfg.get("comune") else None,
                                  cfg.get("comune"), f"({cfg['provincia']})" if cfg.get("provincia") else None]))
    if addr:
        lines.append(addr)
    return lines


# ------------------------------------------------------------
# Endpoints
# ------------------------------------------------------------
@router.post("/presentazioni/genera", status_code=202)
async def presentazione_genera(body: dict, user=Depends(require_user(*STAFF))):
    modalita = body.get("modalita") if body.get("modalita") in ("argomento", "punti") else "argomento"
    contenuto = (body.get("contenuto") or "").strip()
    if not contenuto:
        raise HTTPException(400, "Contenuto obbligatorio")
    stile = body.get("stile") if body.get("stile") in ("ai", "brand") else "brand"
    n_slides = body.get("n_slides")
    n_slides = min(40, max(2, int(n_slides))) if n_slides else None
    con_logo = body.get("con_logo") is not False
    pid = _uid()
    await db.mkt_presentazioni.insert_one({
        "id": pid, "modalita": modalita, "contenuto": contenuto[:2000], "stile": stile,
        "n_slides": n_slides, "stato": "generating", "progress": "Struttura presentazione…",
        "titolo": contenuto[:70], "slides": [], "pptx_path": None, "pdf_path": None,
        "creato_da": user.get("name") or user.get("email"), "created_at": _now_iso(),
    })
    asyncio.create_task(_run_presentazione(pid, modalita, contenuto, n_slides, stile, con_logo))
    return {"id": pid}


async def _run_presentazione(pid: str, modalita: str, contenuto: str,
                             n_slides: int | None, stile: str, con_logo: bool):
    try:
        brand = await _brand_name()
        logo = await _logo_img() if con_logo else None
        contatti = await _linee_contatti()

        vincolo = (f"La presentazione deve avere ESATTAMENTE {n_slides} slide totali."
                   if n_slides else "Scegli tu il numero di slide più adatto (tipicamente 8-14).")
        if modalita == "punti":
            base = ("Impagina in una presentazione professionale i punti chiave forniti dall'utente, "
                    "senza inventare contenuti aggiuntivi: organizza, raggruppa e rifinisci i testi.\n"
                    f"PUNTI CHIAVE DELL'UTENTE:\n{contenuto}")
        else:
            base = f"Crea una presentazione professionale completa sull'argomento: {contenuto}"
        raw = await _llm_text(
            f"{base}\n\nContesto: agenzia assicurativa italiana '{brand}'. {vincolo}\n"
            "Prima slide tipo 'copertina' (titolo max 8 parole + sottotitolo max 15 parole). "
            "Ultima slide tipo 'contatti' (titolo call-to-action, punti = lascia lista vuota). "
            "Slide intermedie tipo 'contenuto': titolo (max 8 parole) + 3-6 punti (max 18 parole l'uno, "
            "italiano perfetto). Ogni slide ha image_prompt (inglese, SOLO sfondo fotografico coerente, no testo).\n"
            'JSON: {"titolo": "titolo presentazione", "slides": [{"tipo": "copertina|contenuto|contatti", '
            '"titolo": "...", "sottotitolo": "...", "punti": ["..."], "image_prompt": "..."}]}')
        piano = _parse_json(raw)
        slides = piano.get("slides") or []
        if n_slides:
            slides = slides[:n_slides]
        if not slides:
            raise ValueError("Piano slide vuoto")
        for sl in slides:
            if sl.get("tipo") not in ("copertina", "contenuto", "contatti"):
                sl["tipo"] = "contenuto"
        slides[0]["tipo"] = "copertina"
        if len(slides) > 1:
            slides[-1]["tipo"] = "contatti"
        for sl in slides:
            if sl["tipo"] == "contatti":
                sl["punti"] = contatti

        # Sfondi AI: copertina dedicata + max 4 sfondi contenuto ciclati
        bg_map: dict[int, bytes] = {}
        if stile == "ai":
            await db.mkt_presentazioni.update_one({"id": pid}, {"$set": {
                "progress": "Generazione sfondi AI…"}})
            content_idx = [i for i, s in enumerate(slides) if s["tipo"] != "copertina"]
            pool: list[bytes] = []
            for k, i in enumerate(content_idx[:4]):
                pool.append(await _genera_sfondo(slides[i].get("image_prompt") or contenuto, orizzontale=True))
                await db.mkt_presentazioni.update_one({"id": pid}, {"$set": {
                    "progress": f"Sfondi AI {k + 2}/{min(4, len(content_idx)) + 1}…"}})
            bg_map[0] = await _genera_sfondo(slides[0].get("image_prompt") or contenuto, orizzontale=True)
            for j, i in enumerate(content_idx):
                bg_map[i] = pool[j % len(pool)] if pool else bg_map[0]

        await db.mkt_presentazioni.update_one({"id": pid}, {"$set": {
            "progress": "Impaginazione slide…"}})
        loop = asyncio.get_event_loop()

        def _render_all():
            bg_only, full = [], []
            for i, sl in enumerate(slides):
                bg = _bg_slide(stile, sl["tipo"], bg_map.get(i), logo)
                out = io.BytesIO()
                bg.save(out, format="JPEG", quality=88)
                bg_only.append(out.getvalue())
                full.append(_draw_slide(bg, sl, stile=stile, n=i + 1, tot=len(slides), brand=brand))
            return bg_only, full

        bg_only, full = await loop.run_in_executor(None, _render_all)

        await db.mkt_presentazioni.update_one({"id": pid}, {"$set": {
            "progress": "Creazione PPTX e PDF…"}})

        def _build_files():
            pptx = _pptx_build(slides, bg_only, stile, brand)
            from PIL import Image
            imgs = [Image.open(io.BytesIO(f)).convert("RGB") for f in full]
            pdf_buf = io.BytesIO()
            imgs[0].save(pdf_buf, format="PDF", save_all=True, append_images=imgs[1:], resolution=144)
            return pptx, pdf_buf.getvalue()

        pptx_data, pdf_data = await loop.run_in_executor(None, _build_files)

        base_path = f"{APP_NAME}/marketing/presentazioni/{pid}"
        slides_meta = []
        for i, jpg in enumerate(full, start=1):
            sp = f"{base_path}/slide_{i}.jpg"
            obj_storage.put_object(sp, jpg, "image/jpeg")
            slides_meta.append({"n": i, "titolo": slides[i - 1].get("titolo"),
                                "tipo": slides[i - 1]["tipo"],
                                "url": f"/api/storage/{sp}", "storage_path": sp})
        pptx_path = f"{base_path}/presentazione.pptx"
        pdf_path = f"{base_path}/presentazione.pdf"
        obj_storage.put_object(pptx_path, pptx_data,
                               "application/vnd.openxmlformats-officedocument.presentationml.presentation")
        obj_storage.put_object(pdf_path, pdf_data, "application/pdf")

        await db.mkt_presentazioni.update_one({"id": pid}, {"$set": {
            "stato": "ready", "titolo": (piano.get("titolo") or contenuto[:70])[:90],
            "slides": slides_meta, "pptx_path": pptx_path, "pdf_path": pdf_path,
            "n_slide_finali": len(slides), "progress": None, "completed_at": _now_iso()}})
    except Exception as e:
        logger.exception("Presentazione %s fallita", pid)
        await db.mkt_presentazioni.update_one({"id": pid}, {"$set": {
            "stato": "errore", "errore": str(e)[:300], "progress": None}})


@router.get("/presentazioni")
async def presentazioni_list(user=Depends(require_user(*STAFF))):
    return await db.mkt_presentazioni.find({}, {"_id": 0}).sort("created_at", -1).to_list(100)


@router.get("/presentazioni/{pid}/file")
async def presentazione_file(pid: str, fmt: str = "pptx", user=Depends(require_user(*STAFF))):
    p = await db.mkt_presentazioni.find_one({"id": pid}, {"_id": 0, "pptx_path": 1, "pdf_path": 1, "titolo": 1})
    if not p:
        raise HTTPException(404, "Presentazione non trovata")
    path = p.get("pdf_path") if fmt == "pdf" else p.get("pptx_path")
    if not path:
        raise HTTPException(404, "File non pronto")
    data, _ = obj_storage.get_object(path)
    media = ("application/pdf" if fmt == "pdf"
             else "application/vnd.openxmlformats-officedocument.presentationml.presentation")
    ext = "pdf" if fmt == "pdf" else "pptx"
    return Response(data, media_type=media,
                    headers={"Content-Disposition": f'attachment; filename="presentazione_{pid[:8]}.{ext}"'})


@router.delete("/presentazioni/{pid}")
async def presentazione_delete(pid: str, user=Depends(require_user(*STAFF))):
    p = await db.mkt_presentazioni.find_one({"id": pid}, {"_id": 0, "slides": 1, "pptx_path": 1, "pdf_path": 1})
    paths = [s.get("storage_path") for s in (p or {}).get("slides") or []]
    paths += [(p or {}).get("pptx_path"), (p or {}).get("pdf_path")]
    for sp in filter(None, paths):
        try:
            obj_storage.delete_object(sp)
        except Exception:
            pass
    await db.mkt_presentazioni.delete_one({"id": pid})
    return {"ok": True}
