"""Rendering professionale del Contratto di Consulenza in PDF e Word (DOCX).

Impagina il testo del contratto (campo `testo` free-editabile) come un vero
documento legale: intestazione con logo agenzia, titolo centrato, premesse,
articoli giustificati (con "Art. N (...)" in grassetto), blocco firme a due
colonne, clausola di approvazione ex art. 1341/1342 c.c. e note legali a piè
di pagina con disclosure sull'uso dell'AI. Lo stesso parsing alimenta sia il
PDF (ReportLab) sia il Word (python-docx) così i modelli restano coerenti.
"""
from __future__ import annotations

import re
from io import BytesIO
from typing import Optional

_ART_RE = re.compile(r"^(Art\.\s*\d+[^—\-–:]*[—\-–:])(.*)$")
_META_PREFIXES = ("Operatore incaricato:", "Corrispettivo pattuito:")
_FIRMA_CLIENTE = "Firma Cliente:"
_FIRMA_CONSULENTE = "Firma Consulente:"

_LEGAL_NOTE = (
    "Documento riservato e confidenziale. Riproduzione vietata senza autorizzazione. "
    "Il presente contratto ha natura di modello e va adattato al caso concreto. "
    "Documento redatto anche con l'ausilio di strumenti di intelligenza artificiale, "
    "verificato dall'intermediario."
)
_APPROVAZIONE = (
    "Ai sensi e per gli effetti degli artt. 1341 e 1342 del Codice Civile, il Cliente "
    "approva specificamente le clausole di cui agli articoli: Art. 4 (Corrispettivo), "
    "Art. 5 (Durata e recesso), Art. 8 (Responsabilità) e Art. 9 (Foro esclusivo)."
)


def _parse(testo: str) -> dict:
    """Divide il testo in: titolo, corpo (lista di (tipo, testo)), firme, meta."""
    lines = (testo or "").split("\n")
    titolo = ""
    body: list[tuple[str, str]] = []
    firma_cliente = ""
    firma_consulente = ""
    meta: list[str] = []
    for raw in lines:
        line = raw.rstrip()
        s = line.strip()
        if not s:
            body.append(("gap", ""))
            continue
        if not titolo and "INCARICO DI CONSULENZA" in s.upper():
            titolo = s
            continue
        if s.startswith(_FIRMA_CLIENTE):
            firma_cliente = s[len(_FIRMA_CLIENTE):].strip()
            continue
        if s.startswith(_FIRMA_CONSULENTE):
            firma_consulente = s[len(_FIRMA_CONSULENTE):].strip()
            continue
        if any(s.startswith(p) for p in _META_PREFIXES):
            meta.append(s)
            continue
        if _ART_RE.match(s):
            body.append(("art", s))
        elif s.lower() in ("tra", "e", "premesso che", "si stipula quanto segue"):
            body.append(("center", s))
        else:
            body.append(("p", s))
    if not titolo:
        titolo = "INCARICO DI CONSULENZA IN MATERIA ASSICURATIVA"
    return {
        "titolo": titolo, "body": body, "meta": meta,
        "firma_cliente": firma_cliente, "firma_consulente": firma_consulente,
    }


# ============================================================
# PDF (ReportLab Platypus)
# ============================================================
def build_pdf(testo: str, *, consulente: Optional[dict] = None,
              cliente_nome: str = "", logo_bytes: Optional[bytes] = None,
              luogo: str = "", data: str = "", firmato: bool = False,
              firmato_il: str = "", firmato_via: str = "") -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.colors import HexColor
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                    TableStyle, Image, HRFlowable)

    cons = consulente or {}
    parsed = _parse(testo)
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=2.6 * cm, bottomMargin=2.2 * cm,
                            leftMargin=2.2 * cm, rightMargin=2.2 * cm,
                            title="Contratto di Consulenza")
    ss = getSampleStyleSheet()
    navy = HexColor("#0f2f4a")
    title = ParagraphStyle("t", parent=ss["Title"], fontName="Helvetica-Bold",
                           fontSize=14, leading=18, alignment=TA_CENTER, textColor=navy,
                           spaceAfter=14)
    body = ParagraphStyle("b", parent=ss["BodyText"], fontName="Helvetica", fontSize=9.5,
                          leading=14, alignment=TA_JUSTIFY, spaceAfter=6)
    center = ParagraphStyle("c", parent=body, alignment=TA_CENTER, fontName="Helvetica-Oblique",
                            spaceBefore=4, spaceAfter=4)
    small = ParagraphStyle("s", parent=ss["BodyText"], fontName="Helvetica", fontSize=7.5,
                           leading=10, alignment=TA_JUSTIFY, textColor=HexColor("#64748b"))
    firm_lbl = ParagraphStyle("fl", parent=ss["BodyText"], fontName="Helvetica-Bold",
                              fontSize=9, leading=12)
    firm_val = ParagraphStyle("fv", parent=ss["BodyText"], fontName="Helvetica", fontSize=8.5,
                              leading=11, textColor=HexColor("#475569"))

    def esc(t: str) -> str:
        return (t or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    story = [Paragraph(esc(parsed["titolo"]), title),
             HRFlowable(width="100%", thickness=1, color=navy, spaceAfter=10)]
    for tipo, txt in parsed["body"]:
        if tipo == "gap":
            story.append(Spacer(1, 4))
        elif tipo == "center":
            story.append(Paragraph(esc(txt), center))
        elif tipo == "art":
            m = _ART_RE.match(txt)
            head, rest = m.group(1), m.group(2)
            story.append(Paragraph(f"<b>{esc(head)}</b>{esc(rest)}", body))
        else:
            story.append(Paragraph(esc(txt), body))

    if parsed["meta"]:
        story.append(Spacer(1, 6))
        for m in parsed["meta"]:
            story.append(Paragraph(esc(m), body))

    # Luogo e data
    story.append(Spacer(1, 10))
    ld = f"Luogo e data: {esc(luogo) or '____________________'}, {esc(data) or '____/____/______'}"
    story.append(Paragraph(ld, body))
    story.append(Spacer(1, 14))

    # Blocco firme a due colonne
    fc = parsed["firma_cliente"] or esc(cliente_nome) or "Il Cliente"
    fk = parsed["firma_consulente"] or esc(cons.get("ragione_sociale", "Il Consulente"))
    sign_note = "Firmato digitalmente via OTP" + (f" ({esc(firmato_via)})" if firmato_via else "")
    cli_cell = [Paragraph("Il Cliente", firm_lbl), Spacer(1, 22),
                HRFlowable(width="80%", thickness=0.6, color=HexColor("#94a3b8")),
                Paragraph(esc(fc), firm_val)]
    con_cell = [Paragraph("Il Consulente", firm_lbl), Spacer(1, 22),
                HRFlowable(width="80%", thickness=0.6, color=HexColor("#94a3b8")),
                Paragraph(esc(fk), firm_val)]
    if firmato:
        cli_cell.append(Paragraph(f"✓ {sign_note} — {esc(firmato_il[:16])}", small))
    t = Table([[cli_cell, con_cell]], colWidths=[doc.width / 2, doc.width / 2])
    t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"),
                           ("LEFTPADDING", (0, 0), (-1, -1), 0),
                           ("RIGHTPADDING", (0, 0), (-1, -1), 14)]))
    story.append(t)

    # Approvazione 1341/1342
    story.append(Spacer(1, 16))
    story.append(Paragraph(_APPROVAZIONE, body))
    story.append(Spacer(1, 8))
    ap = Table([[[Paragraph("Il Cliente (approvazione specifica)", firm_lbl), Spacer(1, 20),
                  HRFlowable(width="70%", thickness=0.6, color=HexColor("#94a3b8"))]]],
               colWidths=[doc.width])
    ap.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0)]))
    story.append(ap)

    def _decorate(canvas, d):
        canvas.saveState()
        W, H = A4
        # Header: logo a destra + ragione sociale consulente a sinistra
        if cons.get("ragione_sociale"):
            canvas.setFont("Helvetica-Bold", 9)
            canvas.setFillColor(navy)
            canvas.drawString(2.2 * cm, H - 1.5 * cm, cons.get("ragione_sociale", "")[:60])
            canvas.setFont("Helvetica", 7)
            canvas.setFillColor(HexColor("#64748b"))
            sub = " · ".join([x for x in [cons.get("partita_iva") and f"P.IVA {cons['partita_iva']}",
                                          cons.get("rui")] if x])
            canvas.drawString(2.2 * cm, H - 1.85 * cm, sub[:90])
        if logo_bytes:
            try:
                from reportlab.lib.utils import ImageReader
                img = ImageReader(BytesIO(logo_bytes))
                iw, ih = img.getSize()
                lw = 3.2 * cm
                lh = lw * ih / iw
                canvas.drawImage(img, W - 2.2 * cm - lw, H - 1.1 * cm - lh, lw, lh,
                                 preserveAspectRatio=True, mask="auto")
            except Exception:
                pass
        canvas.setStrokeColor(HexColor("#cbd5e1"))
        canvas.setLineWidth(0.5)
        canvas.line(2.2 * cm, H - 2.1 * cm, W - 2.2 * cm, H - 2.1 * cm)
        # Footer: nota legale + numero pagina
        canvas.setFont("Helvetica", 6.5)
        canvas.setFillColor(HexColor("#94a3b8"))
        from reportlab.platypus import Paragraph as _P
        note = _P(_LEGAL_NOTE, small)
        note.wrapOn(canvas, W - 4.4 * cm, 2 * cm)
        note.drawOn(canvas, 2.2 * cm, 1.15 * cm)
        canvas.setFont("Helvetica", 7)
        canvas.drawRightString(W - 2.2 * cm, 0.8 * cm, f"Pag. {d.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=_decorate, onLaterPages=_decorate)
    return buf.getvalue()


# ============================================================
# Word (python-docx)
# ============================================================
def build_docx(testo: str, *, consulente: Optional[dict] = None,
               cliente_nome: str = "", logo_bytes: Optional[bytes] = None,
               luogo: str = "", data: str = "") -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    cons = consulente or {}
    parsed = _parse(testo)
    d = Document()
    for section in d.sections:
        section.top_margin = Cm(2.4); section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.2); section.right_margin = Cm(2.2)

    # Intestazione: logo + ragione sociale
    if logo_bytes:
        try:
            hp = d.sections[0].header.paragraphs[0]
            hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            hp.add_run().add_picture(BytesIO(logo_bytes), width=Cm(3.5))
        except Exception:
            pass
    if cons.get("ragione_sociale"):
        h = d.add_paragraph()
        r = h.add_run(cons["ragione_sociale"])
        r.bold = True; r.font.size = Pt(10); r.font.color.rgb = RGBColor(0x0f, 0x2f, 0x4a)
        sub = " · ".join([x for x in [cons.get("partita_iva") and f"P.IVA {cons['partita_iva']}",
                                      cons.get("rui")] if x])
        if sub:
            s = h.add_run("\n" + sub); s.font.size = Pt(7.5); s.font.color.rgb = RGBColor(0x64, 0x74, 0x8b)

    tp = d.add_paragraph(); tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = tp.add_run(parsed["titolo"]); tr.bold = True; tr.font.size = Pt(14)
    tr.font.color.rgb = RGBColor(0x0f, 0x2f, 0x4a)

    for tipo, txt in parsed["body"]:
        if tipo == "gap":
            continue
        p = d.add_paragraph()
        if tipo == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(txt); r.italic = True; r.font.size = Pt(9.5)
        elif tipo == "art":
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            m = _ART_RE.match(txt)
            rb = p.add_run(m.group(1)); rb.bold = True; rb.font.size = Pt(9.5)
            rr = p.add_run(m.group(2)); rr.font.size = Pt(9.5)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            r = p.add_run(txt); r.font.size = Pt(9.5)

    for m in parsed["meta"]:
        r = d.add_paragraph().add_run(m); r.font.size = Pt(9.5)

    d.add_paragraph()
    ld = d.add_paragraph()
    ld.add_run(f"Luogo e data: {luogo or '____________________'}, {data or '____/____/______'}").font.size = Pt(9.5)
    d.add_paragraph()

    # Blocco firme (tabella 2 colonne)
    tbl = d.add_table(rows=1, cols=2)
    fc = parsed["firma_cliente"] or cliente_nome or "Il Cliente"
    fk = parsed["firma_consulente"] or cons.get("ragione_sociale", "Il Consulente")
    for cell, lbl, val in ((tbl.rows[0].cells[0], "Il Cliente", fc),
                           (tbl.rows[0].cells[1], "Il Consulente", fk)):
        rp = cell.paragraphs[0]; rr = rp.add_run(lbl); rr.bold = True; rr.font.size = Pt(9)
        cell.add_paragraph("\n_______________________________")
        vr = cell.add_paragraph().add_run(val); vr.font.size = Pt(8.5); vr.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    d.add_paragraph()
    ap = d.add_paragraph().add_run(_APPROVAZIONE); ap.font.size = Pt(9.5)
    d.add_paragraph("\nIl Cliente (approvazione specifica): _______________________________").runs[0].font.size = Pt(9)

    dn = d.add_paragraph().add_run(_LEGAL_NOTE)
    dn.font.size = Pt(7); dn.font.color.rgb = RGBColor(0x94, 0xa3, 0xb8)

    buf = BytesIO(); d.save(buf); return buf.getvalue()
