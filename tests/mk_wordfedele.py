import os, requests, io
from docx import Document
BASE = [l.split("=",1)[1].strip() for l in open("/app/frontend/.env") if l.startswith("REACT_APP_BACKEND_URL")][0]
s = requests.Session()
r = s.post(f"{BASE}/api/auth/login", json={"email": "admin@assicura.it", "password": "Admin123!"})
print(r.status_code)
tok = r.json().get("access_token") or r.json().get("token")
if tok: s.headers["Authorization"] = f"Bearer {tok}"
d = Document()
d.add_paragraph("Gentile cliente, la presente comunicazione riguarda la sua polizza assicurativa.")
d.add_paragraph("Secondo paragrafo di prova con testo lungo abbastanza.")
buf = io.BytesIO(); d.save(buf); buf.seek(0)
r = s.post(f"{BASE}/api/librerie/modelli/importa-word-fedele", files={"file": ("t.docx", buf, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}, data={"nome": "TEST_WVE_cursor"})
print(r.status_code, r.text[:300])
